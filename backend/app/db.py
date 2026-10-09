"""Meeting dialogue, users, and profiles — in-memory by default; optional PostgreSQL via DATABASE_URL."""
from __future__ import annotations

import io
import threading
import uuid
from datetime import datetime, timezone
from typing import Any

from . import config

_SCHEMA = """
CREATE TABLE IF NOT EXISTS meetings (
    id UUID PRIMARY KEY,
    title TEXT NOT NULL DEFAULT 'Meeting',
    stack TEXT,
    llm TEXT,
    stt TEXT,
    started_at TIMESTAMPTZ NOT NULL DEFAULT now(),
    ended_at TIMESTAMPTZ,
    created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
    user_id UUID
);
"""


class _MemoryStore:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self.users: dict[str, dict[str, Any]] = {}
        self.oauth_index: dict[tuple[str, str], str] = {}
        self.profiles: dict[str, dict[str, Any]] = {}
        self.meetings: dict[str, dict[str, Any]] = {}
        self.lines: dict[str, list[dict[str, Any]]] = {}
        self._line_id = 0

    def clear(self) -> None:
        with self._lock:
            self.users.clear()
            self.oauth_index.clear()
            self.profiles.clear()
            self.meetings.clear()
            self.lines.clear()
            self._line_id = 0


_MEM = _MemoryStore()


def using_postgres() -> bool:
    return bool(config.DATABASE_URL)


def backend() -> str:
    return "postgres" if using_postgres() else "memory"


def reset_memory() -> None:
    """Clear in-memory store (tests / dev only)."""
    if not using_postgres():
        _MEM.clear()


def enabled() -> bool:
    """Storage is always available (in-memory unless DATABASE_URL is set)."""
    return True


def _connect():
    import psycopg
    from psycopg.rows import dict_row

    return psycopg.connect(config.DATABASE_URL, row_factory=dict_row)


def init_db() -> None:
    if not using_postgres():
        return
    with _connect() as conn:
        conn.execute(_SCHEMA)
        conn.execute(
            "ALTER TABLE meetings ADD COLUMN IF NOT EXISTS user_id UUID;"
        )
        conn.commit()


def ping() -> bool:
    if not using_postgres():
        return True
    try:
        with _connect() as conn:
            conn.execute("SELECT 1")
        return True
    except Exception:
        return False


def _ms_to_dt(ms: int) -> datetime:
    return datetime.fromtimestamp(ms / 1000.0, tz=timezone.utc)


def _now() -> datetime:
    return datetime.now(timezone.utc)


def upsert_user(
    provider: str,
    sub: str,
    email: str,
    name: str | None,
    picture: str | None,
) -> str | None:
    if using_postgres():
        return _pg_upsert_user(provider, sub, email, name, picture)
    with _MEM._lock:
        key = (provider, sub)
        if key in _MEM.oauth_index:
            uid = _MEM.oauth_index[key]
            row = _MEM.users[uid]
            row.update({"email": email, "name": name, "picture": picture, "last_login_at": _now()})
            return uid
        uid = str(uuid.uuid4())
        _MEM.oauth_index[key] = uid
        _MEM.users[uid] = {
            "id": uid,
            "email": email,
            "name": name,
            "picture": picture,
            "oauth_provider": provider,
            "oauth_sub": sub,
            "created_at": _now(),
            "last_login_at": _now(),
        }
        return uid


def get_user(user_id: str) -> dict[str, Any] | None:
    if using_postgres():
        with _connect() as conn:
            return conn.execute("SELECT * FROM users WHERE id = %s", (user_id,)).fetchone()
    return _MEM.users.get(user_id)


def get_profile(user_id: str) -> dict[str, Any]:
    if using_postgres():
        with _connect() as conn:
            row = conn.execute("SELECT data FROM user_profiles WHERE user_id = %s", (user_id,)).fetchone()
        if not row:
            return {}
        data = row["data"]
        return dict(data) if isinstance(data, dict) else {}
    return dict(_MEM.profiles.get(user_id, {}))


def save_profile(user_id: str, data: dict[str, Any]) -> None:
    clean = {k: str(v) for k, v in data.items() if v is not None and str(v).strip()}
    if using_postgres():
        import json

        with _connect() as conn:
            conn.execute(
                """INSERT INTO user_profiles (user_id, data, updated_at)
                   VALUES (%s, %s::jsonb, now())
                   ON CONFLICT (user_id) DO UPDATE SET data = EXCLUDED.data, updated_at = now()""",
                (user_id, json.dumps(clean)),
            )
            conn.commit()
        return
    with _MEM._lock:
        _MEM.profiles[user_id] = clean


def meeting_owned_by(meeting_id: str, user_id: str) -> bool:
    if using_postgres():
        with _connect() as conn:
            row = conn.execute(
                "SELECT 1 FROM meetings WHERE id = %s AND user_id = %s",
                (meeting_id, user_id),
            ).fetchone()
        return bool(row)
    m = _MEM.meetings.get(meeting_id)
    return bool(m and m.get("user_id") == user_id)


def create_meeting(
    stack: str | None,
    llm: str | None,
    stt: str | None,
    title: str | None = None,
    user_id: str | None = None,
) -> str | None:
    if using_postgres():
        return _pg_create_meeting(stack, llm, stt, title, user_id)
    mid = str(uuid.uuid4())
    now = _now()
    with _MEM._lock:
        _MEM.meetings[mid] = {
            "id": mid,
            "title": title or "Meeting",
            "stack": stack,
            "llm": llm,
            "stt": stt,
            "user_id": user_id,
            "started_at": now,
            "ended_at": None,
            "line_count": 0,
        }
        _MEM.lines[mid] = []
    return mid


def end_meeting(meeting_id: str) -> None:
    if using_postgres():
        with _connect() as conn:
            conn.execute(
                "UPDATE meetings SET ended_at = now() WHERE id = %s AND ended_at IS NULL",
                (meeting_id,),
            )
            conn.commit()
        return
    with _MEM._lock:
        m = _MEM.meetings.get(meeting_id)
        if m and not m.get("ended_at"):
            m["ended_at"] = _now()


def add_line(meeting_id: str, speaker: str, text: str, spoken_at_ms: int) -> int | None:
    if not text.strip():
        return None
    spoken = _ms_to_dt(spoken_at_ms)
    if using_postgres():
        with _connect() as conn:
            row = conn.execute(
                "SELECT COALESCE(MAX(seq), 0) + 1 AS n FROM dialogue_lines WHERE meeting_id = %s",
                (meeting_id,),
            ).fetchone()
            seq = int(row["n"])
            ins = conn.execute(
                """INSERT INTO dialogue_lines (meeting_id, speaker, text, spoken_at, seq)
                   VALUES (%s, %s, %s, %s, %s) RETURNING id""",
                (meeting_id, speaker.strip() or "Unknown", text.strip(), spoken, seq),
            ).fetchone()
            conn.commit()
            return int(ins["id"])
    with _MEM._lock:
        if meeting_id not in _MEM.meetings:
            return None
        lines = _MEM.lines.setdefault(meeting_id, [])
        seq = len(lines) + 1
        _MEM._line_id += 1
        lid = _MEM._line_id
        lines.append(
            {
                "id": lid,
                "speaker": speaker.strip() or "Unknown",
                "text": text.strip(),
                "spoken_at": spoken,
                "seq": seq,
            }
        )
        _MEM.meetings[meeting_id]["line_count"] = seq
        return lid


def list_meetings(limit: int = 50, user_id: str | None = None) -> list[dict[str, Any]]:
    if using_postgres():
        return _pg_list_meetings(limit, user_id)
    with _MEM._lock:
        rows = list(_MEM.meetings.values())
    if user_id:
        rows = [m for m in rows if m.get("user_id") == user_id]
    else:
        rows = [m for m in rows if not m.get("user_id")]
    rows.sort(key=lambda m: m.get("started_at") or _now(), reverse=True)
    return [_serialize_meeting(m) for m in rows[:limit]]


def get_meeting(meeting_id: str, user_id: str | None = None) -> dict[str, Any] | None:
    if using_postgres():
        return _pg_get_meeting(meeting_id, user_id)
    with _MEM._lock:
        m = _MEM.meetings.get(meeting_id)
        if not m:
            return None
        if user_id and m.get("user_id") != user_id:
            return None
        if user_id is None and m.get("user_id"):
            return None
        lines = list(_MEM.lines.get(meeting_id, []))
    return {**_serialize_meeting(m), "lines": [_serialize_line(x) for x in lines]}


def delete_meeting(meeting_id: str, user_id: str | None = None) -> bool:
    if using_postgres():
        with _connect() as conn:
            if user_id:
                cur = conn.execute("DELETE FROM meetings WHERE id = %s AND user_id = %s", (meeting_id, user_id))
            else:
                cur = conn.execute("DELETE FROM meetings WHERE id = %s AND user_id IS NULL", (meeting_id,))
            conn.commit()
            return cur.rowcount > 0
    with _MEM._lock:
        m = _MEM.meetings.get(meeting_id)
        if not m:
            return False
        if user_id and m.get("user_id") != user_id:
            return False
        if not user_id and m.get("user_id"):
            return False
        del _MEM.meetings[meeting_id]
        _MEM.lines.pop(meeting_id, None)
        return True


def export_docx_bytes(meeting_id: str, user_id: str | None = None) -> bytes | None:
    data = get_meeting(meeting_id, user_id=user_id)
    if not data:
        return None
    import docx

    doc = docx.Document()
    doc.add_heading(data["title"], level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Started: {data['started_at']}\n").italic = True
    if data.get("ended_at"):
        meta.add_run(f"Ended: {data['ended_at']}\n").italic = True
    if data.get("stack"):
        meta.add_run(f"Stack: {data['stack']} ({data.get('llm')}/{data.get('stt')})\n").italic = True
    doc.add_paragraph()
    for line in data.get("lines") or []:
        p = doc.add_paragraph()
        head = p.add_run(f"[{line['time']}] {line['speaker']}: ")
        head.bold = True
        p.add_run(line["text"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pg_upsert_user(provider, sub, email, name, picture) -> str | None:
    uid = str(uuid.uuid4())
    try:
        with _connect() as conn:
            row = conn.execute(
                "SELECT id FROM users WHERE oauth_provider = %s AND oauth_sub = %s",
                (provider, sub),
            ).fetchone()
            if row:
                uid = str(row["id"])
                conn.execute(
                    "UPDATE users SET email = %s, name = %s, picture = %s, last_login_at = now() WHERE id = %s",
                    (email, name, picture, uid),
                )
            else:
                conn.execute(
                    """INSERT INTO users (id, email, name, picture, oauth_provider, oauth_sub)
                       VALUES (%s, %s, %s, %s, %s, %s)""",
                    (uid, email, name, picture, provider, sub),
                )
            conn.commit()
        return uid
    except Exception:
        return None


def _pg_create_meeting(stack, llm, stt, title, user_id) -> str | None:
    try:
        mid = str(uuid.uuid4())
        with _connect() as conn:
            conn.execute(
                "INSERT INTO meetings (id, title, stack, llm, stt, user_id) VALUES (%s, %s, %s, %s, %s, %s)",
                (mid, title or "Meeting", stack, llm, stt, user_id),
            )
            conn.commit()
        return mid
    except Exception:
        return None


def _pg_list_meetings(limit, user_id):
    with _connect() as conn:
        if user_id:
            rows = conn.execute(
                """SELECT m.id, m.title, m.stack, m.llm, m.stt, m.started_at, m.ended_at,
                          COUNT(l.id) AS line_count
                   FROM meetings m
                   LEFT JOIN dialogue_lines l ON l.meeting_id = m.id
                   WHERE m.user_id = %s
                   GROUP BY m.id ORDER BY m.started_at DESC LIMIT %s""",
                (user_id, limit),
            ).fetchall()
        else:
            rows = conn.execute(
                """SELECT m.id, m.title, m.stack, m.llm, m.stt, m.started_at, m.ended_at,
                          COUNT(l.id) AS line_count
                   FROM meetings m
                   LEFT JOIN dialogue_lines l ON l.meeting_id = m.id
                   WHERE m.user_id IS NULL
                   GROUP BY m.id ORDER BY m.started_at DESC LIMIT %s""",
                (limit,),
            ).fetchall()
    return [_serialize_meeting(r) for r in rows]


def _pg_get_meeting(meeting_id, user_id):
    with _connect() as conn:
        if user_id:
            m = conn.execute(
                "SELECT * FROM meetings WHERE id = %s AND user_id = %s",
                (meeting_id, user_id),
            ).fetchone()
        else:
            m = conn.execute("SELECT * FROM meetings WHERE id = %s", (meeting_id,)).fetchone()
        if not m:
            return None
        lines = conn.execute(
            "SELECT speaker, text, spoken_at, seq FROM dialogue_lines WHERE meeting_id = %s ORDER BY seq",
            (meeting_id,),
        ).fetchall()
    return {**_serialize_meeting(m), "lines": [_serialize_line(x) for x in lines]}


def _serialize_meeting(row: dict) -> dict[str, Any]:
    lc = row.get("line_count")
    if lc is None and "id" in row:
        lc = len(_MEM.lines.get(str(row["id"]), [])) if not using_postgres() else 0
    return {
        "id": str(row["id"]),
        "title": row["title"],
        "stack": row.get("stack"),
        "llm": row.get("llm"),
        "stt": row.get("stt"),
        "started_at": _iso(row.get("started_at")),
        "ended_at": _iso(row.get("ended_at")),
        "line_count": int(lc or 0),
    }


def _serialize_line(row: dict) -> dict[str, Any]:
    spoken = row["spoken_at"]
    if isinstance(spoken, (int, float)):
        ms = int(spoken)
        spoken = _ms_to_dt(ms)
    else:
        ms = int(spoken.timestamp() * 1000) if spoken else 0
    return {
        "speaker": row["speaker"],
        "text": row["text"],
        "ts": ms,
        "time": spoken.strftime("%H:%M:%S") if spoken else "",
        "seq": row["seq"],
    }


def _iso(dt) -> str | None:
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc).isoformat()
